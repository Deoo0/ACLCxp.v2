import base64
import io
import tempfile
from pathlib import Path
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from PIL import Image
from rest_framework.test import APIClient
from apps.analytics.ticket_import import template_bytes
from apps.events.tests import make_user, event_data
from apps.events.models import EventCategory, Event
from apps.users.models import IntramuralsTicket
from .models import UploadedImage


class AdminQAChangesTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = make_user("qa-admin", "ADMIN")
        self.student = make_user("qa-student")
        self.client.force_authenticate(self.admin)
        self.category = EventCategory.objects.create(name="QA", slug="qa")

    def workbook(self, numbers):
        workbook = load_workbook(io.BytesIO(template_bytes()))
        for index, number in enumerate(numbers, 2):
            workbook["Tickets"].cell(index, 1, number)
        output = io.BytesIO()
        workbook.save(output)
        return SimpleUploadedFile("tickets.xlsx", output.getvalue())

    def test_excel_template_import_and_leading_zero_numbers(self):
        response = self.client.get("/api/admin/tickets/template/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Tickets", load_workbook(io.BytesIO(response.content)).sheetnames)
        response = self.client.post("/api/admin/tickets/import/", {"file": self.workbook(["000123", "000123456789"])}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual([row["ticket_number"] for row in response.data], ["000123", "000123456789"])
        self.assertEqual(len({row["qr_token"] for row in response.data}), 2)
        self.assertTrue(all(row["status"] == "AVAILABLE" for row in response.data))

    def test_invalid_imports_are_atomic(self):
        IntramuralsTicket.objects.create(ticket_number="111111", qr_token="existing")
        for numbers in [["222222", "111111"], ["222222", "222222"], ["222222", "=1+2"], ["222222", "12345"]]:
            response = self.client.post("/api/admin/tickets/import/", {"file": self.workbook(numbers)}, format="multipart")
            self.assertEqual(response.status_code, 400, response.data)
            self.assertEqual(IntramuralsTicket.objects.count(), 1)
        response = self.client.post("/api/admin/tickets/import/", {"file": SimpleUploadedFile("bad.xlsx", b"invalid")}, format="multipart")
        self.assertEqual(response.status_code, 400)

    def test_ticket_deletion_preserves_redemption_and_is_atomic(self):
        unused = IntramuralsTicket.objects.create(ticket_number="111111", qr_token="unused")
        redeemed = IntramuralsTicket.objects.create(ticket_number="222222", qr_token="redeemed", status="REDEEMED")
        response = self.client.post("/api/admin/tickets/batch-delete/", {"ids": [unused.pk, redeemed.pk]}, format="json")
        self.assertEqual(response.status_code, 409)
        self.assertEqual(IntramuralsTicket.objects.count(), 2)
        self.assertEqual(self.client.delete(f"/api/admin/tickets/{redeemed.pk}/").status_code, 409)
        self.assertEqual(self.client.delete(f"/api/admin/tickets/{unused.pk}/").status_code, 204)
        other = IntramuralsTicket.objects.create(ticket_number="333333", qr_token="other")
        self.assertEqual(self.client.post("/api/admin/tickets/batch-delete/", {"ids": [other.pk]}, format="json").data, {"deleted": 1})

    def test_students_cannot_import_or_delete_tickets(self):
        ticket = IntramuralsTicket.objects.create(ticket_number="111111", qr_token="unused")
        self.client.force_authenticate(self.student)
        for path in ["import", "batch-delete"]:
            self.assertEqual(self.client.post(f"/api/admin/tickets/{path}/", {}, format="json").status_code, 403)
        self.assertEqual(self.client.delete(f"/api/admin/tickets/{ticket.pk}/").status_code, 403)
        self.assertEqual(self.client.get("/api/admin/tickets/template/").status_code, 403)

    def test_text_limits_and_generated_stable_slugs(self):
        data = event_data(self.category)
        data.pop("slug")
        data["description"] = "x" * 3000
        first = self.client.post("/api/events/", data, format="json")
        self.assertEqual(first.status_code, 201, first.data)
        second = self.client.post("/api/events/", data, format="json")
        self.assertEqual(second.status_code, 201, second.data)
        self.assertNotEqual(first.data["slug"], second.data["slug"])
        path = f'/api/events/{first.data["id"]}/'
        changed = self.client.patch(path, {"title": "Renamed", "slug": "manual"}, format="json")
        self.assertEqual(changed.data["slug"], first.data["slug"])
        self.assertEqual(self.client.patch(path, {"description": "x" * 3001}, format="json").status_code, 400)
        self.assertEqual(self.client.post("/api/events/categories/", {"name": "Category without slug"}, format="json").status_code, 201)
        self.assertEqual(self.client.post("/api/events/categories/", {"name": "Too long", "description": "x" * 501}, format="json").status_code, 400)
        limits = self.client.get("/api/admin/field-limits/").data
        self.assertEqual(limits["/events/"]["description"], 3000)
        self.assertEqual(limits["/events/"]["title"], 200)

    @override_settings(STORAGES={"default": {"BACKEND": "apps.core.storage.DatabaseMediaStorage"}})
    def test_upload_survives_instance_disk_change_and_is_readable_across_accounts(self):
        output = io.BytesIO()
        Image.new("RGB", (8, 8), "blue").save(output, format="PNG")
        data = {**event_data(self.category), "banner_image": "data:image/png;base64," + base64.b64encode(output.getvalue()).decode()}
        with tempfile.TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            response = self.client.post("/api/events/", data, format="json")
            self.assertEqual(response.status_code, 201, response.data)
            self.assertEqual(UploadedImage.objects.count(), 1)
            self.assertEqual(list(Path(directory).iterdir()), [])
        with tempfile.TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            for user in [self.admin, self.student, None]:
                self.client.force_authenticate(user)
                photo = self.client.get(response.data["banner_image"])
                self.assertEqual(photo.status_code, 200)
                self.assertEqual(b"".join(photo.streaming_content), output.getvalue())
                self.assertEqual(photo["Cross-Origin-Resource-Policy"], "cross-origin")
                photo.close()

    def test_preserve_existing_uploads_is_repeatable(self):
        name = "events/" + "a" * 32 + ".png"
        data = event_data(self.category)
        data.pop("category")
        Event.objects.create(**data, category=self.category, organizer=self.admin, banner_image=name)
        with tempfile.TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            path = Path(directory) / name
            path.parent.mkdir()
            path.write_bytes(b"existing-image")
            for _ in range(2):
                call_command("preserve_uploads", stdout=io.StringIO())
            self.assertEqual(UploadedImage.objects.count(), 1)
            self.assertEqual(bytes(UploadedImage.objects.get().content), b"existing-image")
