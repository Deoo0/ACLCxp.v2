import io
import re
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from rest_framework.exceptions import ValidationError

MAX_ROWS = 5000


def template_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tickets"
    sheet.append(["ticket_number"])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 28
    for row in range(2, MAX_ROWS + 2):
        sheet.cell(row, 1).number_format = "@"
    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Enter one unique 6- or 12-digit ticket number per row in Tickets."])
    instructions.append(["Keep cells as Text to preserve leading zeros. Do not change the header."])
    instructions.append(["Up to 5,000 tickets. QR tokens are generated automatically on import."])
    instructions.append(["Existing or duplicate ticket numbers reject the entire import."])
    instructions.column_dimensions["A"].width = 100
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def read_ticket_numbers(upload):
    if not upload or not upload.name.lower().endswith(".xlsx"):
        raise ValidationError("Choose an Excel .xlsx file using the ticket template.")
    if upload.size > 5_000_000:
        raise ValidationError("Excel files must be smaller than 5 MB.")
    workbook = None
    try:
        with ZipFile(upload) as archive:
            if sum(item.file_size for item in archive.infolist()) > 25_000_000:
                raise ValidationError("The expanded workbook is too large.")
        upload.seek(0)
        workbook = load_workbook(upload, read_only=True, data_only=False, keep_links=False)
        if "Tickets" not in workbook.sheetnames:
            raise ValidationError("Use the template's Tickets worksheet.")
        sheet = workbook["Tickets"]
        if sheet.max_column and sheet.max_column > 1:
            raise ValidationError("Use only the ticket_number column in Tickets.")
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_col=2)
        header = next(rows)
        if header[0].value != "ticket_number" or header[1].value is not None:
            raise ValidationError("Expected the single column header ticket_number.")
        numbers = []
        for index, row in enumerate(rows, 2):
            cell, extra = row
            if index > MAX_ROWS + 1:
                raise ValidationError("Import at most 5,000 rows.")
            if cell.value is None and extra.value is None:
                continue
            value = str(cell.value).strip()
            if cell.data_type == "f" or extra.value is not None or not re.fullmatch(r"(?:[0-9]{6}|[0-9]{12})", value):
                raise ValidationError(f"Row {index}: enter a 6- or 12-digit ticket number as Text; formulas are not accepted.")
            numbers.append(value)
        if not numbers:
            raise ValidationError("Add at least one ticket number to the Tickets worksheet.")
        if len(set(numbers)) != len(numbers):
            raise ValidationError("Duplicate ticket numbers in the workbook. Nothing was imported.")
        return numbers
    except ValidationError:
        raise
    except (BadZipFile, ValueError, KeyError, TypeError, OSError, EOFError, ParseError, DefusedXmlException) as exc:
        raise ValidationError("The workbook could not be read. Download a fresh template and save it as .xlsx.") from exc
    finally:
        if workbook:
            workbook.close()
